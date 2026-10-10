"""Scanned PDFs, unknown layouts and broken files must end up in review, never as silent bad data."""
import io

from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from extractor import extract_invoice
from extractor.cli import main


def test_scanned_pdf_without_text_goes_to_review(tmp_path):
    img = Image.new("RGB", (800, 1100), "white")
    ImageDraw.Draw(img).text((60, 60), "INVOICE 4471  TOTAL 120.00", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    pdf = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    c.drawImage(ImageReader(buf), 0, 0, *A4)
    c.save()
    inv = extract_invoice(pdf)
    assert inv.layout == "no-text"
    assert inv.needs_review
    assert "OCR" in inv.review_reasons[0]
    assert inv.total is None


def test_corrupt_file_goes_to_review(tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4 this is not really a pdf")
    inv = extract_invoice(bad)
    assert inv.needs_review


def test_unknown_layout_uses_generic_rules_and_asks_for_a_check(tmp_path):
    pdf = tmp_path / "other.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    lines = ["Harbourview Plumbing", "Invoice No: HP-2210", "Date: 2026-05-02",
             "Pipe fitting 4 12.50 50.00", "Callout fee 1 80.00 80.00",
             "Subtotal $130.00", "Tax $13.00", "Total $143.00"]
    for i, line in enumerate(lines):
        c.drawString(72, 780 - i * 20, line)
    c.save()
    inv = extract_invoice(pdf)
    assert inv.layout == "generic"
    assert inv.invoice_number == "HP-2210"
    assert inv.invoice_date == "2026-05-02"
    assert len(inv.line_items) == 2
    assert str(inv.total) == "143.00"
    assert any("unknown layout" in r for r in inv.review_reasons)


def test_cli_writes_all_outputs(tmp_path):
    from tests.conftest import SAMPLES
    out = tmp_path / "out"
    assert main([str(SAMPLES), "--out", str(out)]) == 0
    for name in ("invoices.xlsx", "invoices.csv", "line_items.csv", "invoices.json"):
        assert (out / name).stat().st_size > 0
    from openpyxl import load_workbook
    wb = load_workbook(out / "invoices.xlsx")
    assert wb.sheetnames == ["Invoices", "Line Items", "Review"]
    assert wb["Invoices"].max_row == 13
    assert wb["Review"].max_row >= 2
