"""Excel (3 sheets), CSV and JSON writers."""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3B63")
HEADER_FONT = Font(bold=True, color="FFFFFF")
REVIEW_FILL = PatternFill("solid", fgColor="FDE7E7")
MONEY_FMT = {"USD": '"$"#,##0.00', "EUR": '#,##0.00 "EUR"', "GBP": '"GBP "#,##0.00'}

INVOICE_COLS = ["source_file", "layout", "invoice_number", "vendor", "invoice_date", "currency",
                "subtotal", "tax", "total", "line_count", "status"]
ITEM_COLS = ["source_file", "invoice_number", "currency", "description", "quantity", "unit_price", "amount"]


def invoice_row(inv):
    return [inv.source_file, inv.layout, inv.invoice_number, inv.vendor, inv.invoice_date, inv.currency,
            inv.subtotal, inv.tax, inv.total, len(inv.line_items),
            "NEEDS REVIEW" if inv.needs_review else "OK"]


def item_rows(inv):
    for li in inv.line_items:
        yield [inv.source_file, inv.invoice_number, inv.currency, li.description, li.quantity,
               li.unit_price, li.amount]


def _sheet(ws, cols, rows, money_cols=(), currency_col=None):
    ws.append([c.replace("_", " ").title() for c in cols])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([float(v) if hasattr(v, "quantize") else v for v in row])
        r = ws.max_row
        cur = ws.cell(r, cols.index(currency_col) + 1).value if currency_col else None
        for name in money_cols:
            ws.cell(r, cols.index(name) + 1).number_format = MONEY_FMT.get(cur, "#,##0.00")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(cols, 1):
        width = max(len(str(c.value or "")) for c in ws[get_column_letter(i)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)


def write_excel(invoices, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    _sheet(ws, INVOICE_COLS, (invoice_row(i) for i in invoices), ("subtotal", "tax", "total"), "currency")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, len(INVOICE_COLS)).value == "NEEDS REVIEW":
            for c in ws[r]:
                c.fill = REVIEW_FILL
    _sheet(wb.create_sheet("Line Items"), ITEM_COLS, (row for i in invoices for row in item_rows(i)),
           ("unit_price", "amount"), "currency")
    review = wb.create_sheet("Review")
    rows = [[i.source_file, i.invoice_number, reason] for i in invoices for reason in i.review_reasons]
    _sheet(review, ["source_file", "invoice_number", "reason"], rows or [["-", "-", "nothing to review"]])
    wb.save(path)


def write_csv(invoices, out_dir: Path):
    with open(out_dir / "invoices.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(INVOICE_COLS)
        w.writerows(invoice_row(i) for i in invoices)
    with open(out_dir / "line_items.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(ITEM_COLS)
        w.writerows(row for i in invoices for row in item_rows(i))


def write_json(invoices, path: Path):
    path.write_text(json.dumps([i.to_dict() for i in invoices], indent=2), encoding="utf-8")
