"""Render the README screenshots from real output (needs playwright + chromium)."""
import base64
import html
import io
import json
from pathlib import Path

import pypdfium2 as pdfium
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots"
BASE_CSS = """*{box-sizing:border-box;margin:0;padding:0}
body{width:1440px;height:900px;background:#0d1321;font-family:'Inter','Segoe UI',system-ui,sans-serif;color:#e6ebf5;overflow:hidden}
.mono{font-family:'JetBrains Mono','DejaVu Sans Mono',monospace}"""


def page_png(pdf_path, scale=1.6):
    page = pdfium.PdfDocument(str(pdf_path))[0]
    img = page.render(scale=scale).to_pil()
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode()


def colour_json(obj):
    s = html.escape(json.dumps(obj, indent=2))
    import re
    s = re.sub(r"(&quot;[a-z_]+&quot;)(:)", r'<span style="color:#7aa2ff">\1</span>\2', s)
    s = re.sub(r": (&quot;.*?&quot;)", r': <span style="color:#7ee0a1">\1</span>', s)
    s = re.sub(r": (true|false|null)", r': <span style="color:#ffb86b">\1</span>', s)
    return s


def shot_side_by_side():
    inv = json.loads((ROOT / "sample_output" / "invoices.json").read_text())[0]
    inv["line_items"] = inv["line_items"][:3] + ([{"...": f"{len(inv['line_items']) - 3} more"}] if len(inv["line_items"]) > 3 else [])
    for k in ("source_file", "layout"):
        inv.pop(k)
    img = page_png(ROOT / "samples" / inv_file())
    return f"""<html><head><style>{BASE_CSS}
.wrap{{display:grid;grid-template-columns:600px 1fr;gap:28px;padding:28px 36px}}
h2{{font-size:15px;letter-spacing:2px;color:#8a95ad;margin-bottom:12px;text-transform:uppercase}}
.pdf{{background:#fff;border-radius:10px;overflow:hidden;height:800px}}
.pdf img{{width:100%;display:block;margin-top:-10px}}
pre{{background:#111a2e;border:1px solid #24304b;border-radius:10px;padding:18px 22px;font-size:13.5px;line-height:1.5;height:800px;overflow:hidden}}
</style></head><body><div class="wrap">
<div><h2>Input &middot; samples/01_northwind.pdf</h2><div class="pdf"><img src="data:image/png;base64,{img}"></div></div>
<div><h2>Output &middot; sample_output/invoices.json (first record)</h2><pre class="mono">{colour_json(inv)}</pre></div>
</div></body></html>"""


def inv_file():
    return "01_northwind.pdf"


def sheet_html(ws, max_rows=14):
    rows = list(ws.iter_rows(values_only=True))[:max_rows]
    out = ["<table>"]
    money = {j for j, h in enumerate(rows[0]) if h in ("Subtotal", "Tax", "Total", "Unit Price", "Amount")}
    for i, row in enumerate(rows):
        review = i > 0 and "NEEDS REVIEW" in [str(v) for v in row]
        cls = ' class="rev"' if review else ""
        tag = "th" if i == 0 else "td"
        cells = []
        for j, v in enumerate(row):
            if i > 0 and j in money and isinstance(v, (int, float)):
                v = f"{v:,.2f}"
            cells.append(f"<{tag}>{html.escape(str(v if v is not None else ''))}</{tag}>")
        out.append(f"<tr{cls}>{''.join(cells)}</tr>")
    out.append("</table>")
    return "".join(out)


def shot_excel():
    wb = load_workbook(ROOT / "sample_output" / "invoices.xlsx")
    inv, review = wb["Invoices"], wb["Review"]
    return f"""<html><head><style>{BASE_CSS}
body{{background:#e9edf3;color:#1d2433}}
.bar{{background:#1d6f42;color:#fff;padding:12px 24px;font-size:15px;font-weight:600}}
.wrap{{padding:18px 24px}}
table{{border-collapse:collapse;font-size:12.5px;background:#fff;width:100%}}
th{{background:#1f3b63;color:#fff;text-align:left;padding:7px 8px;font-weight:600}}
td{{border:1px solid #dde3ec;padding:6px 8px;white-space:nowrap}}
tr.rev td{{background:#fde7e7}}
h3{{font-size:13px;margin:16px 0 8px;color:#45506a;letter-spacing:1px;text-transform:uppercase}}
.tabs{{display:flex;gap:2px;margin-top:10px}}
.tabs span{{background:#cfd6e2;padding:6px 16px;font-size:12px;border-radius:0 0 6px 6px}}
.tabs span.on{{background:#fff;font-weight:700;color:#1d6f42}}
</style></head><body><div class="bar">sample_output/invoices.xlsx</div><div class="wrap">
<h3>Invoices sheet</h3>{sheet_html(inv, 13)}
<div class="tabs"><span class="on">Invoices</span><span>Line Items</span><span>Review</span></div>
<h3>Review sheet</h3>{sheet_html(review, 5)}
</div></body></html>"""


def shot_terminal():
    run = (ROOT / "docs" / "run_log.txt").read_text().splitlines()
    acc = (ROOT / "docs" / "accuracy_log.txt").read_text().splitlines()[-2:]
    tests = (ROOT / "docs" / "pytest_q.txt").read_text().splitlines()[-1:]

    def fmt(line):
        e = html.escape(line)
        for word, col in (("REVIEW", "#ff8f8f"), ("WARNING", "#ffb86b"), ("INFO", "#5fd0a0"), ("passed", "#5fd0a0")):
            e = e.replace(word, f'<span style="color:{col}">{word}</span>')
        return e
    body = ['<span class="p">$</span> python -m extractor samples/ --out sample_output/']
    body += [fmt(l) for l in run]
    body += ["", '<span class="p">$</span> python scripts/check_accuracy.py | tail -2'] + [fmt(l) for l in acc]
    body += ["", '<span class="p">$</span> python -m pytest -q'] + [fmt(l) for l in tests]
    return f"""<html><head><style>{BASE_CSS}
.win{{margin:22px 30px;background:#0a0f1a;border:1px solid #24304b;border-radius:12px;overflow:hidden;height:856px}}
.t{{background:#141c2e;padding:11px 16px;display:flex;gap:8px;align-items:center}}
.t i{{width:12px;height:12px;border-radius:50%;display:block}}
.t span{{margin-left:12px;color:#8a95ad;font-size:13px}}
pre{{padding:16px 20px;font-size:13px;line-height:1.55;color:#cdd6e6}}
.p{{color:#7aa2ff}}
</style></head><body><div class="win"><div class="t"><i style="background:#ff5f56"></i><i style="background:#ffbd2e"></i><i style="background:#27c93f"></i><span class="mono">pdf-invoice-extractor - bash</span></div>
<pre class="mono">{chr(10).join(body)}</pre></div></body></html>"""


def main():
    SHOTS.mkdir(parents=True, exist_ok=True)
    jobs = [("01_pdf_to_json.png", shot_side_by_side()), ("02_excel_output.png", shot_excel()),
            ("03_cli_and_tests.png", shot_terminal())]
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={"width": 1440, "height": 900}, device_scale_factor=1)
        for name, doc in jobs:
            pg.set_content(doc)
            pg.wait_for_timeout(300)
            pg.screenshot(path=str(SHOTS / name))
            print(SHOTS / name)
        b.close()


if __name__ == "__main__":
    main()
